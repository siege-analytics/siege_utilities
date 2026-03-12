"""
Unit tests for the siege_utilities reporting module.

Covers ReportGenerator, ChartGenerator, LegendManager, ClientBrandingManager,
PowerPointGenerator, and image_utils across 35+ tests.
"""

import base64
import io
import logging
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import MagicMock, patch

import pandas as pd
import pytest

# ---------------------------------------------------------------------------
# Optional dependency availability flags
# ---------------------------------------------------------------------------
try:
    import reportlab  # noqa: F401
    from reportlab.platypus import Image as RLImage
    from reportlab.lib.units import inch

    REPORTLAB_AVAILABLE = True
except ImportError:
    REPORTLAB_AVAILABLE = False

try:
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    MATPLOTLIB_AVAILABLE = True
except ImportError:
    MATPLOTLIB_AVAILABLE = False

try:
    from pptx import Presentation  # noqa: F401

    PPTX_AVAILABLE = True
except ImportError:
    PPTX_AVAILABLE = False

try:
    import folium  # noqa: F401

    FOLIUM_AVAILABLE = True
except ImportError:
    FOLIUM_AVAILABLE = False

# Lazy-safe imports — modules themselves guard optional deps internally
from siege_utilities.reporting.client_branding import ClientBrandingManager
from siege_utilities.reporting.legend_manager import (
    ColorScheme,
    LegendManager,
    LegendPosition,
)
from siege_utilities.reporting.image_utils import decode_rl_image, save_rl_image

# These need reportlab at import time
pytestmark_reportlab = pytest.mark.skipif(
    not REPORTLAB_AVAILABLE, reason="reportlab not installed"
)


# ============================================================================
# FIXTURES
# ============================================================================

@pytest.fixture
def tmp_output(tmp_path):
    """Temporary output directory for generators."""
    d = tmp_path / "output"
    d.mkdir()
    return d


@pytest.fixture
def branding_manager(tmp_path):
    """ClientBrandingManager with isolated config dir."""
    return ClientBrandingManager(config_dir=tmp_path / "branding")


@pytest.fixture
def legend_manager():
    return LegendManager()


@pytest.fixture
def sample_branding_config():
    return {
        "name": "Test Client",
        "colors": {
            "primary": "#112233",
            "secondary": "#445566",
            "text_color": "#000000",
        },
        "fonts": {"default_font": "Helvetica"},
    }


def _make_rl_image_stub(b64_png: str = None):
    """Return a mock with a base64 data-URI filename, mimicking ChartGenerator output."""
    if b64_png is None:
        # 1x1 transparent PNG
        b64_png = (
            "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR4"
            "2mNkYPhfDwAChwGA60e6kgAAAABJRU5ErkJggg=="
        )
    stub = SimpleNamespace(filename=f"data:image/png;base64,{b64_png}")
    return stub


# ============================================================================
# 1. ReportGenerator tests
# ============================================================================

@pytestmark_reportlab
class TestReportGeneratorImport:
    """Basic import and instantiation tests for ReportGenerator."""

    def test_import(self):
        from siege_utilities.reporting.report_generator import ReportGenerator
        assert ReportGenerator is not None

    def test_instantiation(self, tmp_output):
        from siege_utilities.reporting.report_generator import ReportGenerator
        rg = ReportGenerator("siege_analytics", output_dir=tmp_output)
        assert rg.client_name == "siege_analytics"
        assert rg.output_dir == tmp_output

    def test_instantiation_unknown_client_falls_back(self, tmp_output):
        from siege_utilities.reporting.report_generator import ReportGenerator
        rg = ReportGenerator("nonexistent_client_xyz", output_dir=tmp_output)
        assert rg.branding_config == {}

    def test_add_section(self, tmp_output):
        from siege_utilities.reporting.report_generator import ReportGenerator
        rg = ReportGenerator("siege_analytics", output_dir=tmp_output)
        report = {"sections": []}
        updated = rg.add_section(report, "text", "Intro", "Hello world")
        assert len(updated["sections"]) == 1
        assert updated["sections"][0]["title"] == "Intro"
        assert updated["sections"][0]["type"] == "text"

    def test_add_text_section(self, tmp_output):
        from siege_utilities.reporting.report_generator import ReportGenerator
        rg = ReportGenerator("siege_analytics", output_dir=tmp_output)
        report = {"sections": []}
        updated = rg.add_text_section(report, "Summary", "Some text here")
        assert updated["sections"][0]["content"] == "Some text here"

    def test_add_table_section_with_dataframe(self, tmp_output):
        from siege_utilities.reporting.report_generator import ReportGenerator
        rg = ReportGenerator("siege_analytics", output_dir=tmp_output)
        report = {"sections": []}
        df = pd.DataFrame({"A": [1, 2], "B": [3, 4]})
        updated = rg.add_table_section(report, "Data Table", df)
        section = updated["sections"][0]
        assert section["type"] == "table"
        # Headers extracted from DataFrame columns
        assert section["content"]["headers"] == ["A", "B"]
        # Data should include header row + data rows
        assert len(section["content"]["data"]) == 3  # header + 2 rows

    def test_add_table_section_with_list(self, tmp_output):
        from siege_utilities.reporting.report_generator import ReportGenerator
        rg = ReportGenerator("siege_analytics", output_dir=tmp_output)
        report = {"sections": []}
        data = [[1, 2], [3, 4]]
        updated = rg.add_table_section(report, "Numbers", data, headers=["X", "Y"])
        section = updated["sections"][0]
        assert section["content"]["data"][0] == ["X", "Y"]

    def test_create_analytics_report(self, tmp_output):
        from siege_utilities.reporting.report_generator import ReportGenerator
        rg = ReportGenerator("siege_analytics", output_dir=tmp_output)
        result = rg.create_analytics_report(
            title="Test Report",
            charts=[],
            data_summary="Testing data",
            insights=["Insight 1"],
            recommendations=["Rec 1"],
            executive_summary="Exec summary",
            methodology="Our methodology",
        )
        assert result["title"] == "Test Report"
        assert result["type"] == "analytics_report"
        # Should have sections for exec summary, methodology, data summary, insights, recs
        section_types = [s["type"] for s in result["sections"]]
        assert "executive_summary" in section_types
        assert "methodology" in section_types
        assert "insights" in section_types
        assert "recommendations" in section_types

    def test_create_comprehensive_report(self, tmp_output):
        from siege_utilities.reporting.report_generator import ReportGenerator
        rg = ReportGenerator("siege_analytics", output_dir=tmp_output)
        sections = [{"type": "text", "title": "Chapter 1", "content": "Body text"}]
        result = rg.create_comprehensive_report(
            title="Full Report",
            author="Tester",
            sections=sections,
            table_of_contents=True,
        )
        assert result["title"] == "Full Report"
        assert result["metadata"]["author"] == "Tester"
        # Should have title_page + toc + our section
        section_types = [s["type"] for s in result["sections"]]
        assert "title_page" in section_types
        assert "table_of_contents" in section_types

    def test_add_section_creates_sections_key(self, tmp_output):
        from siege_utilities.reporting.report_generator import ReportGenerator
        rg = ReportGenerator("siege_analytics", output_dir=tmp_output)
        report = {}  # no 'sections' key yet
        updated = rg.add_section(report, "text", "Title", "Content")
        assert "sections" in updated
        assert len(updated["sections"]) == 1


# ============================================================================
# 2. ChartGenerator tests
# ============================================================================

@pytestmark_reportlab
class TestChartGenerator:
    """Tests for ChartGenerator methods."""

    def test_import(self):
        from siege_utilities.reporting.chart_generator import ChartGenerator
        assert ChartGenerator is not None

    def test_instantiation_defaults(self):
        from siege_utilities.reporting.chart_generator import ChartGenerator
        cg = ChartGenerator()
        assert cg.output_dir == Path.home() / ".siege_utilities"
        assert cg.max_chart_width is None
        assert cg.max_chart_height is None

    def test_instantiation_with_config(self, tmp_output):
        from siege_utilities.reporting.chart_generator import ChartGenerator
        config = {"colors": {"primary": "#FF0000", "secondary": "#00FF00"}}
        cg = ChartGenerator(branding_config=config, output_dir=tmp_output,
                            max_chart_width=6.0, max_chart_height=8.5)
        assert cg.default_colors["primary"] == "#FF0000"
        assert cg.max_chart_width == 6.0

    def test_scale_dimensions_no_limits(self):
        from siege_utilities.reporting.chart_generator import ChartGenerator
        cg = ChartGenerator()
        w, h = cg._scale_dimensions(10.0, 8.0)
        assert w == 10.0
        assert h == 8.0

    def test_scale_dimensions_with_limits(self):
        from siege_utilities.reporting.chart_generator import ChartGenerator
        cg = ChartGenerator(max_chart_width=6.0, max_chart_height=4.0)
        w, h = cg._scale_dimensions(12.0, 8.0)
        # Scale factor should be min(6/12, 4/8) = 0.5
        assert abs(w - 6.0) < 0.01
        assert abs(h - 4.0) < 0.01

    def test_scale_dimensions_preserves_aspect_ratio(self):
        from siege_utilities.reporting.chart_generator import ChartGenerator
        cg = ChartGenerator(max_chart_width=5.0)
        w, h = cg._scale_dimensions(10.0, 6.0)
        assert abs(w - 5.0) < 0.01
        assert abs(h - 3.0) < 0.01  # 6 * (5/10)

    @pytest.mark.skipif(not MATPLOTLIB_AVAILABLE, reason="matplotlib not installed")
    def test_save_figure_as_vector_svg(self, tmp_output):
        from siege_utilities.reporting.chart_generator import ChartGenerator
        cg = ChartGenerator(output_dir=tmp_output)
        fig, ax = plt.subplots()
        ax.plot([1, 2, 3], [1, 4, 9])
        out = cg.save_figure_as_vector(fig, tmp_output / "test_chart", fmt="svg")
        assert out is not None
        assert out.suffix == ".svg"
        assert out.exists()
        plt.close(fig)

    @pytest.mark.skipif(not MATPLOTLIB_AVAILABLE, reason="matplotlib not installed")
    def test_save_figure_as_vector_eps(self, tmp_output):
        from siege_utilities.reporting.chart_generator import ChartGenerator
        cg = ChartGenerator(output_dir=tmp_output)
        fig, ax = plt.subplots()
        ax.plot([1, 2], [3, 4])
        out = cg.save_figure_as_vector(fig, tmp_output / "test_chart", fmt="eps")
        assert out is not None
        assert out.suffix == ".eps"
        assert out.exists()
        plt.close(fig)

    @pytest.mark.skipif(not MATPLOTLIB_AVAILABLE, reason="matplotlib not installed")
    def test_save_figure_as_vector_pdf(self, tmp_output):
        from siege_utilities.reporting.chart_generator import ChartGenerator
        cg = ChartGenerator(output_dir=tmp_output)
        fig, ax = plt.subplots()
        ax.plot([1, 2], [3, 4])
        out = cg.save_figure_as_vector(fig, tmp_output / "test_chart", fmt="pdf")
        assert out is not None
        assert out.suffix == ".pdf"
        assert out.exists()
        plt.close(fig)

    def test_save_figure_as_vector_invalid_format(self, tmp_output):
        from siege_utilities.reporting.chart_generator import ChartGenerator
        cg = ChartGenerator(output_dir=tmp_output)
        fig = MagicMock()
        result = cg.save_figure_as_vector(fig, tmp_output / "test_chart", fmt="bmp")
        assert result is None

    @pytest.mark.skipif(not MATPLOTLIB_AVAILABLE, reason="matplotlib not installed")
    def test_matplotlib_to_reportlab_image(self, tmp_output):
        from siege_utilities.reporting.chart_generator import ChartGenerator
        cg = ChartGenerator(output_dir=tmp_output)
        fig, ax = plt.subplots(figsize=(4, 3))
        ax.bar(["A", "B"], [10, 20])
        rl_img = cg._matplotlib_to_reportlab_image(fig, 4.0, 3.0)
        assert rl_img is not None
        # Should be a ReportLab Image with a data URI
        assert hasattr(rl_img, "filename") or hasattr(rl_img, "_filename")

    @pytest.mark.skipif(not MATPLOTLIB_AVAILABLE, reason="matplotlib not installed")
    def test_create_bar_chart(self, tmp_output):
        from siege_utilities.reporting.chart_generator import ChartGenerator
        cg = ChartGenerator(output_dir=tmp_output)
        df = pd.DataFrame({"category": ["A", "B", "C"], "value": [10, 20, 30]})
        img = cg.create_bar_chart(df, x_column="category", y_column="value",
                                  title="Test Bar")
        assert img is not None

    @pytest.mark.skipif(not MATPLOTLIB_AVAILABLE, reason="matplotlib not installed")
    def test_create_placeholder_chart(self, tmp_output):
        from siege_utilities.reporting.chart_generator import ChartGenerator
        cg = ChartGenerator(output_dir=tmp_output)
        img = cg._create_placeholder_chart(4.0, 3.0, "Placeholder text")
        assert img is not None

    def test_default_dpi(self):
        from siege_utilities.reporting.chart_generator import ChartGenerator
        cg = ChartGenerator()
        assert cg.default_dpi == 150


# ============================================================================
# 3. LegendManager tests
# ============================================================================

class TestLegendManager:
    """Tests for LegendManager."""

    def test_import(self):
        from siege_utilities.reporting.legend_manager import LegendManager
        assert LegendManager is not None

    def test_instantiation(self, legend_manager):
        assert legend_manager.branding_config == {}
        assert len(legend_manager.color_schemes) == 6

    def test_instantiation_with_branding(self):
        config = {"colors": {"primary": "#AABB00"}}
        lm = LegendManager(branding_config=config)
        assert lm.branding_config == config

    def test_color_for_intensity_zero(self, legend_manager):
        color = legend_manager.get_color_for_intensity(0.0, ColorScheme.BLUE)
        # intensity 0 should return the low color
        assert color.lower() == "#87ceeb"

    def test_color_for_intensity_one(self, legend_manager):
        color = legend_manager.get_color_for_intensity(1.0, ColorScheme.BLUE)
        assert color.lower() == "#1e3a5f"

    def test_color_for_intensity_midpoint(self, legend_manager):
        color = legend_manager.get_color_for_intensity(0.5, ColorScheme.BLUE)
        # Should be between low and high — just check it's a valid hex
        assert color.startswith("#")
        assert len(color) == 7

    def test_color_schemes_all_present(self, legend_manager):
        for scheme in ColorScheme:
            assert scheme in legend_manager.color_schemes
            info = legend_manager.color_schemes[scheme]
            assert "low" in info
            assert "high" in info
            assert "name" in info

    def test_legend_position_enum(self):
        assert LegendPosition.BEST.value == "best"
        assert LegendPosition.OUTSIDE.value == "outside"

    def test_get_legend_position_for_pie(self, legend_manager):
        pos = legend_manager.get_legend_position_for_chart_type("pie", 3)
        assert pos == LegendPosition.RIGHT

    def test_get_legend_position_for_pie_many(self, legend_manager):
        pos = legend_manager.get_legend_position_for_chart_type("pie", 10)
        assert pos == LegendPosition.BOTTOM

    def test_should_show_legend_pie(self, legend_manager):
        assert legend_manager.should_show_legend("pie", 5) is True

    def test_should_show_legend_bar_single(self, legend_manager):
        # Single series, few data points — no legend needed
        assert legend_manager.should_show_legend("bar", 3, series_count=1) is False

    def test_should_show_legend_bar_multi(self, legend_manager):
        assert legend_manager.should_show_legend("bar", 3, series_count=2) is True

    @pytest.mark.skipif(not REPORTLAB_AVAILABLE, reason="reportlab not installed")
    def test_create_heatmap_legend_table(self, legend_manager):
        table = legend_manager.create_heatmap_legend_table(0, 100, "Count")
        assert table is not None

    @pytest.mark.skipif(not MATPLOTLIB_AVAILABLE, reason="matplotlib not installed")
    def test_create_heatmap_legend_elements(self, legend_manager):
        elements = legend_manager.create_heatmap_legend_elements(0, 100)
        assert len(elements) == 5  # default levels

    def test_default_legend_styles(self, legend_manager):
        styles = legend_manager.legend_styles
        assert styles["frameon"] is True
        assert styles["fontsize"] == 9

    def test_get_legend_position_scatter(self, legend_manager):
        pos = legend_manager.get_legend_position_for_chart_type("scatter", 5)
        assert pos == LegendPosition.BEST


# ============================================================================
# 4. ClientBrandingManager tests
# ============================================================================

class TestClientBrandingManager:
    """Tests for ClientBrandingManager."""

    def test_import(self):
        from siege_utilities.reporting.client_branding import ClientBrandingManager
        assert ClientBrandingManager is not None

    def test_instantiation(self, branding_manager):
        assert branding_manager.config_dir.exists()

    def test_get_builtin_branding_siege(self, branding_manager):
        config = branding_manager.get_client_branding("siege_analytics")
        assert config is not None
        assert config["name"] == "Siege Analytics"
        assert "colors" in config
        assert "fonts" in config

    def test_get_builtin_branding_masai(self, branding_manager):
        config = branding_manager.get_client_branding("masai_interactive")
        assert config is not None
        assert config["name"] == "Masai Interactive"

    def test_get_builtin_branding_hillcrest(self, branding_manager):
        config = branding_manager.get_client_branding("hillcrest")
        assert config is not None

    def test_get_unknown_client_returns_none(self, branding_manager):
        result = branding_manager.get_client_branding("totally_unknown_client")
        assert result is None

    def test_list_clients_includes_builtins(self, branding_manager):
        clients = branding_manager.list_clients()
        # Builtins are stored lowercase
        assert "siege_analytics" in clients
        assert "masai_interactive" in clients
        assert "hillcrest" in clients

    def test_create_client_branding(self, branding_manager, sample_branding_config):
        path = branding_manager.create_client_branding("test_client", sample_branding_config)
        assert path.exists()
        assert path.suffix == ".yaml"

    def test_create_then_get(self, branding_manager, sample_branding_config):
        branding_manager.create_client_branding("new_client", sample_branding_config)
        config = branding_manager.get_client_branding("new_client")
        assert config is not None
        assert config["name"] == "Test Client"

    def test_create_missing_required_field_raises(self, branding_manager):
        bad_config = {"name": "Incomplete"}  # missing 'colors' and 'fonts'
        with pytest.raises(ValueError, match="Missing required field"):
            branding_manager.create_client_branding("bad", bad_config)

    def test_validate_branding_config_valid(self, branding_manager, sample_branding_config):
        errors = branding_manager.validate_branding_config(sample_branding_config)
        assert len(errors) == 0

    def test_validate_branding_config_missing_fields(self, branding_manager):
        errors = branding_manager.validate_branding_config({})
        assert len(errors) >= 3  # name, colors, fonts

    def test_validate_branding_config_missing_colors(self, branding_manager):
        config = {"name": "X", "colors": {}, "fonts": {"default_font": "Helvetica"}}
        errors = branding_manager.validate_branding_config(config)
        assert any("primary" in e for e in errors)

    def test_delete_predefined_template_fails(self, branding_manager):
        result = branding_manager.delete_client_branding("siege_analytics")
        assert result is False

    def test_branding_summary(self, branding_manager):
        summary = branding_manager.get_branding_summary("siege_analytics")
        assert summary["client_name"] == "Siege Analytics"
        assert summary["has_logo"] is True
        assert "primary" in summary["colors"]

    def test_export_branding_yaml(self, branding_manager, tmp_path):
        export_path = tmp_path / "export.yaml"
        result = branding_manager.export_branding_config("siege_analytics", export_path)
        assert result is True
        assert export_path.exists()

    def test_export_branding_json(self, branding_manager, tmp_path):
        export_path = tmp_path / "export.json"
        result = branding_manager.export_branding_config("siege_analytics", export_path)
        assert result is True
        assert export_path.exists()

    def test_color_extraction_from_branding(self, branding_manager):
        config = branding_manager.get_client_branding("siege_analytics")
        colors = config["colors"]
        assert colors["primary"].startswith("#")
        assert colors["secondary"].startswith("#")
        assert colors["accent"].startswith("#")


# ============================================================================
# 5. PowerPointGenerator tests
# ============================================================================

@pytest.mark.skipif(not PPTX_AVAILABLE, reason="python-pptx not installed")
class TestPowerPointGenerator:
    """Tests for PowerPointGenerator."""

    def test_import(self):
        from siege_utilities.reporting.powerpoint_generator import PowerPointGenerator
        assert PowerPointGenerator is not None

    def test_instantiation(self, tmp_output):
        from siege_utilities.reporting.powerpoint_generator import PowerPointGenerator
        pg = PowerPointGenerator("test_client", output_dir=tmp_output)
        assert pg.client_name == "test_client"
        assert pg.output_dir == tmp_output

    def test_create_analytics_presentation(self, tmp_output):
        from siege_utilities.reporting.powerpoint_generator import PowerPointGenerator
        pg = PowerPointGenerator("test_client", output_dir=tmp_output)
        report_data = {
            "executive_summary": "This is a summary.",
            "metrics": {"visits": 1000, "conversions": 50},
        }
        path = pg.create_analytics_presentation(report_data, "Test Pres")
        assert path.exists()
        assert path.suffix == ".pptx"

    def test_create_custom_presentation(self, tmp_output):
        from siege_utilities.reporting.powerpoint_generator import PowerPointGenerator
        pg = PowerPointGenerator("test_client", output_dir=tmp_output)
        config = {
            "title": "Custom Presentation",
            "type": "demo",
            "slides": [
                {"type": "content", "title": "Slide 1", "content": "Hello"},
            ],
        }
        path = pg.create_custom_presentation(config)
        assert path.exists()

    def test_add_slide_section(self, tmp_output):
        from siege_utilities.reporting.powerpoint_generator import PowerPointGenerator
        pg = PowerPointGenerator("test_client", output_dir=tmp_output)
        pres = {"sections": []}
        updated = pg.add_slide_section(pres, "text_slide", "Title", "Content")
        assert len(updated["sections"]) == 1
        assert updated["sections"][0]["title"] == "Title"

    def test_add_text_slide(self, tmp_output):
        from siege_utilities.reporting.powerpoint_generator import PowerPointGenerator
        pg = PowerPointGenerator("test_client", output_dir=tmp_output)
        pres = {"sections": []}
        updated = pg.add_text_slide(pres, "Info Slide", "Some information here")
        assert updated["sections"][0]["type"] == "text_slide"
        assert updated["sections"][0]["content"]["text"] == "Some information here"

    def test_comprehensive_presentation_structure(self, tmp_output):
        from siege_utilities.reporting.powerpoint_generator import PowerPointGenerator
        pg = PowerPointGenerator("test_client", output_dir=tmp_output)
        result = pg.create_comprehensive_presentation(
            title="Big Deck",
            author="Tester",
            sections=[{"type": "content", "title": "S1", "content": "Body"}],
        )
        assert result["title"] == "Big Deck"
        assert result["type"] == "comprehensive_presentation"
        assert result["metadata"]["author"] == "Tester"


@pytest.mark.skipif(PPTX_AVAILABLE, reason="test only when python-pptx is missing")
class TestPowerPointGeneratorMissing:
    """Verify graceful handling when python-pptx is not installed."""

    def test_import_error_on_instantiation(self, tmp_output):
        from siege_utilities.reporting.powerpoint_generator import PowerPointGenerator
        with pytest.raises(ImportError):
            PowerPointGenerator("test_client", output_dir=tmp_output)


# ============================================================================
# 6. image_utils tests
# ============================================================================

class TestImageUtils:
    """Tests for decode_rl_image and save_rl_image."""

    def test_decode_rl_image_valid(self):
        stub = _make_rl_image_stub()
        result = decode_rl_image(stub)
        assert result is not None
        assert isinstance(result, bytes)
        # Should be valid PNG header
        assert result[:4] == b"\x89PNG"

    def test_decode_rl_image_none_for_non_data_uri(self):
        stub = SimpleNamespace(filename="/some/path.png")
        result = decode_rl_image(stub)
        assert result is None

    def test_decode_rl_image_none_for_no_filename(self):
        stub = SimpleNamespace(value=42)  # no filename attribute
        result = decode_rl_image(stub)
        assert result is None

    def test_decode_rl_image_none_for_none_input(self):
        result = decode_rl_image(None)
        assert result is None

    def test_save_rl_image_writes_file(self, tmp_path):
        stub = _make_rl_image_stub()
        out_path = tmp_path / "saved.png"
        result = save_rl_image(stub, out_path)
        assert result == out_path
        assert out_path.exists()
        content = out_path.read_bytes()
        assert content[:4] == b"\x89PNG"

    def test_save_rl_image_creates_parent_dirs(self, tmp_path):
        stub = _make_rl_image_stub()
        out_path = tmp_path / "nested" / "dir" / "image.png"
        result = save_rl_image(stub, out_path)
        assert out_path.exists()

    def test_save_rl_image_non_data_uri(self, tmp_path):
        stub = SimpleNamespace(filename="/some/regular/path.png")
        out_path = tmp_path / "nope.png"
        result = save_rl_image(stub, out_path)
        # Should return path but not write since decoding fails
        assert result == out_path
        assert not out_path.exists()

    def test_decode_rl_image_roundtrip(self):
        """Encode known bytes, wrap in stub, decode, verify match."""
        original = b"\x89PNG\r\n\x1a\nfake_png_data"
        b64 = base64.b64encode(original).decode()
        stub = SimpleNamespace(filename=f"data:image/png;base64,{b64}")
        decoded = decode_rl_image(stub)
        assert decoded == original


# ============================================================================
# 7. Module-level convenience functions
# ============================================================================

class TestConvenienceFunctions:
    """Test module-level convenience functions from legend_manager."""

    def test_get_optimal_legend_position(self):
        from siege_utilities.reporting.legend_manager import get_optimal_legend_position
        pos = get_optimal_legend_position("pie", 3)
        assert pos == "right"

    def test_get_optimal_legend_position_bar(self):
        from siege_utilities.reporting.legend_manager import get_optimal_legend_position
        pos = get_optimal_legend_position("bar", 10)
        assert pos == "outside"

    @pytest.mark.skipif(not REPORTLAB_AVAILABLE, reason="reportlab not installed")
    def test_create_heatmap_legend_table_convenience(self):
        from siege_utilities.reporting.legend_manager import create_heatmap_legend_table
        table = create_heatmap_legend_table(0, 1000, "Donations", "green")
        assert table is not None
